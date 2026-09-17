"""Export the offload/resident matrix JSON evidence into the comparison workbook.

The JSON schema comes from ``benchmark_fsa_matrix.py``: each arm directory holds
``server_info.json`` plus ``summary.json`` / ``final.json`` (with a ``cases`` list)
and a ``<mode>_<length>/summary.json`` file for every finished case.  A case that was
never reported stays ``待测试``; an explicit ``no_data`` / failed case becomes ``无数据``.
Latencies use equal-weight request means when raw request records are available.
Throughput uses the measured batch makespan; missing measurements are never synthesized.
"""

import argparse
import json
import re
import shutil
import sys
import tempfile
from copy import copy
from datetime import datetime
from pathlib import Path

from fsa_matrix_metrics import request_latency_means

try:
    import openpyxl
except ModuleNotFoundError:  # pragma: no cover - depends on the checkout layout
    _ANALYSIS_TOOLS = Path(__file__).resolve().parents[2] / "results" / "analysis_tools"
    if _ANALYSIS_TOOLS.is_dir():
        sys.path.insert(0, str(_ANALYSIS_TOOLS))
    import openpyxl

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font

ARMS = ("baseline", "offload")
ARM_LABELS = {"baseline": "基线", "offload": "offload"}
MODES = ("single", "saturation")
MODE_ORDER = ("saturation", "single")
MODE_LABELS = {"single": "单请求", "saturation": "64请求批次饱和"}
DEFAULT_LENGTHS = (128000, 256000, 512000, 1000000)

DECIMAL_K = 1000
OUTPUT_TOKENS = 128
CONTEXT_LENGTH = 1048576
SERVICE_REQUEST_CAP = 64

OK = "ok"
PENDING = "待测试"
NO_DATA = "无数据"
STATUS_LABELS = {OK: "成功", NO_DATA: NO_DATA, PENDING: PENDING}

DEPLOYMENT = "混部：单服务集成 P/D"
PARALLEL = "TP8"
WEIGHTS = "GLM-5.3-Flash / FP8 权重 / BF16 KV"

MATRIX_SHEET = "Sheet1"
SINGLE_SHEET = "单请求性能"
CONFIG_SHEET = "测试配置"
DETAIL_SHEET = "明细与状态"
GAIN_SHEET = "收益对比"
GENERATED_SHEETS = (SINGLE_SHEET, CONFIG_SHEET, DETAIL_SHEET, GAIN_SHEET)

COLUMNS = tuple("ABCDEFGHIJKLMN")
HEADER_ROWS = (1, 2)
FIRST_DATA_ROW = 3
HEADER_MERGES = ("A1:A2", "B1:D1", "E1:N1")

METRIC_COLUMNS = (
    ("ttft_ms", "TTFT", "ms", "latency"),
    ("itl_ms", "ITL", "ms", "latency"),
    ("tpot_ms", "TPOT", "ms", "latency"),
    ("output_tokens_per_s", "输出吞吐", "token/s", "throughput"),
    ("output_tokens_per_s_per_card", "单卡输出吞吐", "token/s", "throughput"),
    ("e2e_tokens_per_s", "E2E吞吐", "token/s", "throughput"),
    ("e2e_tokens_per_s_per_card", "单卡E2E吞吐", "token/s", "throughput"),
)

SERVER_INFO_FIELDS = ("context_length", "max_running_requests", "max_total_tokens", "max_req_input_len", "page_size",
                      "max_total_num_tokens", "tp_size", "disable_radix_cache",
                      "speculative_algorithm", "model_path")

CASE_DIRECTORY = re.compile(r"^(?P<mode>single|saturation)_(?P<length>\d+)$")
BACKUP_LABEL = "备份文件"
BACKUP_STAMP = "%Y%m%d_%H%M%S"

CENTER = Alignment(horizontal="center", vertical="center")
BOLD = Font(bold=True)

BATCH_NOTES = (
    "注：本表为 64 请求并发工作负载（批次饱和 / batch saturation）的吞吐结果，并非无限稳态；"
    "客户端并发与服务上限均为 64（max_running_requests=64），实际并发以“明细与状态”中的"
    " actual_max_decode_batch 为准。",
    "口径：上下文 1048576 tokens 所有用例共用；1K=1000 tokens（decimal）；输出 128 tokens=0.128K；"
    "TTFT/ITL/TPOT 为均值（ms）；吞吐为 token/s；E2E 吞吐包含输入+输出；未启用 profiling。",
    "延迟口径：TTFT_i=服务端prefill_finished_time-forward_entry_time，排除首次调度前排队；ITL_i=请求i相邻token间隔均值；"
    "TPOT_i=(请求i末token到达-首token到达)/(输出token数-1)。三项均先逐请求计算，再对请求等权平均。",
    "整批耗时仅用于吞吐：所有请求输出token总数/整批耗时；请求吞吐=完成请求数/整批耗时。"
    "缺少服务端时间戳的TTFT标为未采集推理起点，不能用含排队的客户端TTFT替代。"
    "ITL/TPOT从首token之后计时，保留生成期间的调度暂停，不使用整批耗时计算。",
)

SINGLE_NOTES = (
    "注：本表为单请求（client_concurrency=1）结果，用于与 Sheet1 的 64 请求批次饱和结果对照。",
    "口径与 Sheet1 一致：上下文 1048576 tokens 共用；1K=1000 tokens；输出 128 tokens=0.128K；"
    "无 profiling。",
)

DETAIL_HEADER = ["模式", "上下文", "输入/k", "方案", "状态", "客户端并发", "提交请求数",
                 "实际最大decode batch", "实际最大full tokens", "TTFT均值(ms)", "ITL均值(ms)",
                 "TPOT均值(ms)", "输出吞吐(token/s)", "E2E吞吐(token/s)", "备注/原因", "证据文件"]

DETAIL_NOTES = (
    "注：client_concurrency 为客户端并发；64 请求为批次饱和（batch saturation）工作负载，不是无限稳态。"
    "服务并发上限 max_running_requests=64，实际并发以 actual_max_decode_batch 为准。",
    "成功用例的 TTFT/ITL/TPOT 为均值（ms），吞吐为 token/s；待测试 = 该用例尚无结果文件；"
    "无数据 = 用例显式返回 no_data/failed 且绝不合成指标。",
    "TTFT/ITL/TPOT均为逐请求指标的等权平均；批量总耗时只用于吞吐计算。",
)

GAIN_HEADER = ["模式", "上下文", "输入/k", "指标", "单位", "基线", "offload",
               "增益%（offload 相对基线，正值为更优）"]

GAIN_NOTES = (
    "注：仅当同一长度下基线与 offload 均为成功用例时才计算增益；缺失或 no_data 的组合不做任何推算。",
    "增益%（延迟）= (基线 - offload) / 基线；增益%（吞吐）= (offload - 基线) / 基线。",
)


def _read_json(path):
    path = Path(path)
    if not path.is_file():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _case_key(mode, length):
    if mode not in MODES or length is None:
        return None
    try:
        return (mode, int(length))
    except (TypeError, ValueError):
        return None


def _cases_from_payload(payload):
    if not isinstance(payload, dict) or not isinstance(payload.get("cases"), list):
        return {}
    cases = {}
    for case in payload["cases"]:
        if not isinstance(case, dict):
            continue
        key = _case_key(case.get("mode"), case.get("input_tokens"))
        if key is not None:
            cases[key] = case
    return cases


def load_arm(directory):
    """Collect one arm's cases from its summary/final file and its per-case dirs."""
    directory = Path(directory)
    arm = {"directory": directory, "server_info": {}, "cases": {}, "evidence": {},
           "sources": [], "meta": {}}
    if not directory.is_dir():
        return arm
    arm["server_info"] = _read_json(directory / "server_info.json") or {}
    for name in ("summary.json", "final.json"):
        payload = _read_json(directory / name)
        if payload is None:
            continue
        arm["sources"].append(directory / name)
        if isinstance(payload, dict):
            arm["meta"] = payload
        for key, case in _cases_from_payload(payload).items():
            arm["cases"].setdefault(key, case)
            arm["evidence"].setdefault(key, directory / name)
    for entry in sorted(directory.iterdir()):
        match = CASE_DIRECTORY.match(entry.name) if entry.is_dir() else None
        if match is None:
            continue
        case = _read_json(entry / "summary.json")
        if case is None:
            continue
        key = _case_key(match.group("mode"), match.group("length"))
        arm["cases"][key] = case          # the per-case file is the freshest copy
        arm["evidence"][key] = entry / "summary.json"
    return arm


def _metrics(record):
    if not isinstance(record, dict):
        return {}
    metrics = record.get("metrics")
    if not isinstance(metrics, dict):
        return {}
    requests = record.get("requests")
    if requests and all(all(field in request for field in ("ttft_s", "tpot_ms", "itl_ms"))
                        for request in requests):
        return {**metrics, **request_latency_means(requests)}
    return metrics


def case_status(record):
    """OK only for a reported success with metrics; missing cases stay pending."""
    if not isinstance(record, dict) or record.get("status") is None:
        return PENDING
    if record.get("status") != OK:
        return NO_DATA
    return OK if _metrics(record) else NO_DATA


def _number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _metric_value(record, field):
    """Measured value, or the pending / no-data marker instead of an invented number."""
    state = case_status(record)
    if state != OK:
        return state
    value = _metrics(record).get(field)
    if field == "ttft_ms" and value is None:
        return "未采集推理起点"
    return value if _number(value) else PENDING


def _field_value(record, field):
    state = case_status(record)
    if state != OK:
        return state
    value = record.get(field)
    return value if _number(value) else PENDING


def _record_note(record):
    if not isinstance(record, dict):
        return "没有该用例的结果文件（待测试）"
    if case_status(record) == OK:
        parts = []
        resident = record.get("estimated_resident_requests")
        if _number(resident):
            parts.append(f"估算常驻请求={resident}")
        errors = record.get("errors")
        if isinstance(errors, list) and errors:
            parts.append(f"errors={len(errors)}")
        return "；".join(parts)
    reason = record.get("reason") or record.get("error") or record.get("status")
    errors = record.get("errors")
    if isinstance(errors, list) and errors:
        first = errors[0]
        detail = first.get("error") if isinstance(first, dict) else first
        return f"status={record.get('status')}；reason={reason}；首个错误={detail}"
    return f"status={record.get('status')}；reason={reason}"


def _evidence(arm, key):
    return str(arm["evidence"].get(key, arm["directory"]))


def _write_matrix_row(sheet, row_index, arm_name, length, record):
    values = [
        ARM_LABELS[arm_name],
        CONTEXT_LENGTH,
        length / DECIMAL_K,
        OUTPUT_TOKENS / DECIMAL_K,
        DEPLOYMENT,
        PARALLEL,
        WEIGHTS,
    ]
    values.extend(_metric_value(record, field) for field, _label, _unit, _kind in METRIC_COLUMNS)
    for offset, value in enumerate(values):
        cell = sheet.cell(row=row_index, column=1 + offset, value=value)
        cell.alignment = CENTER


def _clear_values(sheet, first_row, last_column):
    last_row = max(sheet.max_row, first_row)
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, max_col=last_column):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def _write_matrix(sheet, arms, mode, lengths):
    """Overwrite the generated data area below the preserved template headers."""
    _clear_values(sheet, FIRST_DATA_ROW, len(COLUMNS))
    row_index = FIRST_DATA_ROW
    for length in lengths:
        for arm_name in ARMS:
            _write_matrix_row(sheet, row_index, arm_name, length,
                              arms[arm_name]["cases"].get((mode, length)))
            row_index += 1
    return row_index


def _write_note(sheet, row_index, lines):
    for offset, line in enumerate(lines):
        cell = sheet.cell(row=row_index + offset, column=1, value=line)
        cell.alignment = Alignment(horizontal="left")


def _copy_header(template, sheet):
    for row in HEADER_ROWS:
        for column in range(1, len(COLUMNS) + 1):
            source = template.cell(row=row, column=column)
            target = sheet.cell(row=row, column=column)
            target.value = source.value
            target.font = copy(source.font)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
    for merged in HEADER_MERGES:
        sheet.merge_cells(merged)


def _copy_layout(template, sheet):
    for letter, dimension in template.column_dimensions.items():
        if dimension.width:
            sheet.column_dimensions[letter].width = dimension.width


def _fresh_sheet(workbook, title, last_column):
    if title in workbook.sheetnames:
        sheet = workbook[title]
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
    else:
        sheet = workbook.create_sheet(title)
    _clear_values(sheet, 1, last_column)
    return sheet


def _backup_marker(workbook):
    if CONFIG_SHEET not in workbook.sheetnames:
        return None
    sheet = workbook[CONFIG_SHEET]
    for row in sheet.iter_rows(min_row=1, max_row=max(sheet.max_row, 1), max_col=2):
        if len(row) > 1 and row[0].value == BACKUP_LABEL and row[1].value:
            return str(row[1].value)
    return None


def ensure_backup(workbook, workbook_path, timestamp=None):
    """Create the one-time timestamped backup; later exports reuse the existing copy."""
    recorded = _backup_marker(workbook)
    if recorded and Path(recorded).is_file():
        return recorded, False
    pattern = f"{workbook_path.stem}.backup-*{workbook_path.suffix}"
    existing = sorted(workbook_path.parent.glob(pattern))
    if existing:
        return str(existing[-1]), False
    stamp = timestamp or datetime.now().strftime(BACKUP_STAMP)
    backup = workbook_path.with_name(f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}")
    shutil.copy2(workbook_path, backup)
    return str(backup), True


def _text(value):
    if value is None:
        return "（缺失）"
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _arm_info_row(arms, arm_name):
    arm = arms[arm_name]
    info = arm["server_info"]
    meta = arm["meta"] if isinstance(arm["meta"], dict) else {}
    sources = "；".join(str(path) for path in arm["sources"]) or "（缺失）"
    return ([ARM_LABELS[arm_name], str(arm["directory"]), sources, _text(meta.get("host"))]
            + [_text(info.get(field)) for field in SERVER_INFO_FIELDS])


def _write_config(sheet, arms, backup, stamp):
    _clear_values(sheet, 1, 20)
    lines = [
        ("导出时间", stamp),
        ("备份文件", backup),
        ("基线目录", str(arms["baseline"]["directory"])),
        ("offload目录", str(arms["offload"]["directory"])),
        ("结果状态（基线）", _text(arms["baseline"]["meta"].get("status")
                                   if isinstance(arms["baseline"]["meta"], dict) else None)),
        ("结果状态（offload）", _text(arms["offload"]["meta"].get("status")
                                     if isinstance(arms["offload"]["meta"], dict) else None)),
        ("", ""),
        ("通用口径", ""),
        ("十进制约定", f"1K = {DECIMAL_K} tokens（decimal K）；输出 {OUTPUT_TOKENS} tokens = "
                       f"{OUTPUT_TOKENS / DECIMAL_K:.3f}K"),
        ("上下文长度", f"{CONTEXT_LENGTH} tokens，所有用例共用"),
        ("部署方案", DEPLOYMENT),
        ("并行策略", PARALLEL),
        ("权重/精度", WEIGHTS),
        ("投机采样", "MTP0（未启用 MTP / speculative）"),
        ("服务并发上限", f"max_running_requests = {SERVICE_REQUEST_CAP}"),
        ("工作负载", "64 请求 = 批次饱和（batch saturation），并非无限稳态；实际并发以"
                     " actual_max_decode_batch 为准，不等于 64 的服务上限"),
        ("延迟口径", "TTFT / ITL / TPOT 主单元格为均值，单位 ms"),
        ("吞吐口径", "吞吐单位 token/s；E2E 吞吐 = (输入 tokens + 输出 tokens) / 耗时"),
        ("Profiling", "未启用 profiling"),
        ("状态约定", f"{PENDING} = 该用例尚无结果文件；{NO_DATA} = 用例显式返回 no_data/failed，"
                     "不合成任何成功指标"),
    ]
    row_index = 1
    for key, value in lines:
        sheet.cell(row=row_index, column=1, value=key).font = BOLD
        sheet.cell(row=row_index, column=2, value=value)
        row_index += 1
    row_index += 1
    sheet.cell(row=row_index, column=1,
               value="服务与容量（来自 server_info.json / final.json）").font = BOLD
    row_index += 1
    header = ["方案", "证据目录", "结果文件", "host"] + list(SERVER_INFO_FIELDS)
    for offset, title in enumerate(header):
        sheet.cell(row=row_index, column=1 + offset, value=title).font = BOLD
    row_index += 1
    for arm_name in ARMS:
        for offset, value in enumerate(_arm_info_row(arms, arm_name)):
            sheet.cell(row=row_index, column=1 + offset, value=value)
        row_index += 1
    return row_index


def _detail_row(arms, arm_name, mode, length):
    arm = arms[arm_name]
    record = arm["cases"].get((mode, length))
    state = case_status(record)
    return [
        MODE_LABELS[mode],
        CONTEXT_LENGTH,
        length / DECIMAL_K,
        ARM_LABELS[arm_name],
        STATUS_LABELS[state],
        _field_value(record, "client_concurrency"),
        _field_value(record, "submitted_requests"),
        _field_value(record, "actual_max_decode_batch"),
        _field_value(record, "actual_max_full_tokens"),
        _metric_value(record, "ttft_ms"),
        _metric_value(record, "itl_ms"),
        _metric_value(record, "tpot_ms"),
        _metric_value(record, "output_tokens_per_s"),
        _metric_value(record, "e2e_tokens_per_s"),
        _record_note(record),
        _evidence(arm, (mode, length)),
    ]


def _write_details(sheet, arms, lengths):
    _clear_values(sheet, 1, 24)
    for offset, title in enumerate(DETAIL_HEADER):
        sheet.cell(row=1, column=1 + offset, value=title).font = BOLD
    row_index = 2
    for mode in MODE_ORDER:
        for length in lengths:
            for arm_name in ARMS:
                for offset, value in enumerate(_detail_row(arms, arm_name, mode, length)):
                    sheet.cell(row=row_index, column=1 + offset, value=value)
                row_index += 1
    row_index += 1
    for line in DETAIL_NOTES:
        sheet.cell(row=row_index, column=1, value=line)
        row_index += 1
    return row_index
def _gain(kind, baseline, offload):
    if not _number(baseline) or not _number(offload) or baseline == 0:
        return None
    if kind == "latency":
        return (baseline - offload) / baseline * 100
    return (offload - baseline) / baseline * 100


def _write_gains(sheet, arms, lengths):
    _clear_values(sheet, 1, 24)
    for offset, title in enumerate(GAIN_HEADER):
        sheet.cell(row=1, column=1 + offset, value=title).font = BOLD
    row_index = 2
    pairs = 0
    skipped = []
    for mode in MODE_ORDER:
        for length in lengths:
            baseline = arms["baseline"]["cases"].get((mode, length))
            offload = arms["offload"]["cases"].get((mode, length))
            if case_status(baseline) != OK or case_status(offload) != OK:
                skipped.append(
                    f"{MODE_LABELS[mode]} {length / DECIMAL_K:g}K："
                    f"基线={STATUS_LABELS[case_status(baseline)]}，"
                    f"offload={STATUS_LABELS[case_status(offload)]}")
                continue
            pairs += 1
            for field, label, unit, kind in METRIC_COLUMNS:
                baseline_value = _metrics(baseline).get(field)
                offload_value = _metrics(offload).get(field)
                gain = _gain(kind, baseline_value, offload_value)
                values = [MODE_LABELS[mode], CONTEXT_LENGTH, length / DECIMAL_K, label, unit,
                          baseline_value if _number(baseline_value) else PENDING,
                          offload_value if _number(offload_value) else PENDING,
                          gain if gain is not None else "无法计算"]
                for offset, value in enumerate(values):
                    sheet.cell(row=row_index, column=1 + offset, value=value)
                row_index += 1
    row_index += 1
    if not pairs:
        sheet.cell(row=row_index, column=1, value="暂无可计算收益的成对成功结果").font = BOLD
        row_index += 1
    if skipped:
        sheet.cell(row=row_index, column=1, value="未成对（不计算收益）").font = BOLD
        row_index += 1
        for line in skipped:
            sheet.cell(row=row_index, column=1, value=line)
            row_index += 1
    for line in GAIN_NOTES:
        sheet.cell(row=row_index, column=1, value=line)
        row_index += 1
    return row_index


def _lengths(arms):
    lengths = set(DEFAULT_LENGTHS)
    for arm in arms.values():
        lengths.update(length for _mode, length in arm["cases"])
    return sorted(lengths)


def export_workbook(workbook_path, baseline_dir, offload_dir, timestamp=None):
    """Populate the generated areas of the workbook and report what was written."""
    workbook_path = Path(workbook_path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"找不到工作簿：{workbook_path}")
    workbook = openpyxl.load_workbook(workbook_path)
    if MATRIX_SHEET not in workbook.sheetnames:
        raise ValueError(f"工作簿缺少模板工作表 {MATRIX_SHEET}")
    stamp = timestamp or datetime.now().strftime(BACKUP_STAMP)
    backup, created = ensure_backup(workbook, workbook_path, stamp)
    arms = {"baseline": load_arm(baseline_dir), "offload": load_arm(offload_dir)}
    lengths = _lengths(arms)
    matrix = workbook[MATRIX_SHEET]
    single = _fresh_sheet(workbook, SINGLE_SHEET, len(COLUMNS))
    _copy_header(matrix, single)
    _copy_layout(matrix, single)
    _write_note(matrix, _write_matrix(matrix, arms, "saturation", lengths) + 1, BATCH_NOTES)
    _write_note(single, _write_matrix(single, arms, "single", lengths) + 1, SINGLE_NOTES)
    _write_config(_fresh_sheet(workbook, CONFIG_SHEET, 24), arms, backup, stamp)
    _write_details(_fresh_sheet(workbook, DETAIL_SHEET, 24), arms, lengths)
    _write_gains(_fresh_sheet(workbook, GAIN_SHEET, 24), arms, lengths)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=workbook_path.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    actual_path = workbook_path
    try:
        workbook.save(temporary_path)
        try:
            temporary_path.replace(workbook_path)
        except PermissionError:
            actual_path = workbook_path.with_name(workbook_path.stem + "_测试结果" + workbook_path.suffix)
            temporary_path.replace(actual_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return {"workbook": str(actual_path), "requested_workbook": str(workbook_path), "backup": backup, "backup_created": created,
            "generated_at": stamp, "lengths": list(lengths),
            "sheets": list(workbook.sheetnames)}


def main(argv=None):
    parser = argparse.ArgumentParser(description="导出 offload 性能对比工作簿")
    parser.add_argument("--workbook", type=Path, required=True, help="待填充的 xlsx 模板")
    parser.add_argument("--offload", type=Path, required=True, help="offload 臂结果目录")
    parser.add_argument("--baseline", type=Path, required=True, help="基线臂结果目录")
    args = parser.parse_args(argv)
    report = export_workbook(args.workbook, args.baseline, args.offload)
    try:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    except UnicodeEncodeError:  # pragma: no cover - legacy console codepage
        print(json.dumps(report, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
